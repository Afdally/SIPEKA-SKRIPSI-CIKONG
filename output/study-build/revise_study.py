from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

root = Path(__file__).resolve().parents[2]
source = root / 'output/Panduan_Pertanyaan_Sidang_SIPEKA.docx'
output = root / 'output/Panduan_Pertanyaan_Sidang_SIPEKA_Revisi.docx'
doc = Document(source)

replacements = {
'- unused -': '',
'Alasan penelitian membahas arsitektur microservices dan monolith.': 'Alasan penerapan microservices pada SIPEKA dan kedudukan monolith sebagai baseline evaluasi.',
'Indikator yang digunakan untuk membandingkannya.': 'Indikator evaluasi yang sesuai dengan tujuan dan rumusan masalah penelitian.',
'“Penelitian ini membangun SIPEKA untuk mengintegrasikan pelaporan masyarakat dan penanganan kasus oleh petugas. Selain menghasilkan sistem, penelitian membandingkan performa arsitektur microservices dan monolith menggunakan skenario beban yang sama.”': '“Fokus penelitian saya adalah penerapan arsitektur microservices pada SIPEKA untuk mengintegrasikan pelaporan masyarakat dan penanganan kasus. Monolith digunakan sebagai baseline, yaitu acuan untuk mengevaluasi performa implementasi pada skenario pengujian yang ditentukan.”',
'Kontribusi penelitian: hasil perbandingan performa kedua arsitektur dalam konteks SIPEKA.': 'Kontribusi penelitian: rancangan pembagian layanan dan integrasi data pada SIPEKA, disertai evaluasi implementasi dengan monolith sebagai baseline. Klaim kebaruan harus dijelaskan terhadap penelitian terdahulu.',
'Kemampuan sistem menghadapi pertambahan pengguna/request.': 'Kinerja implementasi pada beban yang diuji. Jangan menyimpulkan kemampuan scaling jika eksperimen scaling belum dilakukan.',
'2 Metodologi pengujian': '2 Metodologi evaluasi implementasi',
'12 Pertanyaan prioritas latihan': '12 Jawaban sepuluh pertanyaan prioritas',
'Mengapa hasil microservices bisa lebih lambat?': 'Mengapa hasil microservices bisa lebih lambat atau hanya berbeda sedikit?',
}
for p in doc.paragraphs:
    if p.text in replacements:
        was_bold=bool(p.runs and p.runs[0].bold)
        p.text=replacements[p.text]
        if was_bold:
            for r in p.runs: r.bold=True

def before(anchor,text,bold=False):
    p=anchor.insert_paragraph_before(text)
    p.paragraph_format.keep_with_next=bold
    if bold:
        for r in p.runs:r.bold=True
    return p

anchor=next(p for p in doc.paragraphs if p.text=='1 Pertanyaan pembuka penelitian')
before(anchor,'Kedudukan penelitian dan pengujian',True)
before(anchor,'Gunakan kerangka penerapan microservices sebagai fokus dan monolith sebagai baseline hanya jika sesuai dengan judul, rumusan masalah, tujuan, dan metodologi skripsi. Pengujian tetap memiliki unsur komparatif. Istilah baseline menjelaskan kedudukannya sebagai acuan, bukan menghilangkan fakta bahwa dua implementasi dibandingkan.')
before(anchor,'Jika ditanya mengapa judulnya bukan perbandingan arsitektur',True)
before(anchor,'“Objek utama penelitian saya adalah penerapan microservices pada SIPEKA. Implementasi monolith digunakan untuk memberi acuan dalam menilai hasil pengujian. Karena itu judul menekankan penerapan, sedangkan pengujian komparatif ditempatkan sebagai bagian evaluasinya.”')
before(anchor,'Bukti yang perlu ditunjukkan: keterkaitan judul dengan rumusan masalah, rancangan pembagian service, implementasi komunikasi antarservice, serta hasil pengujian. Jika seluruh rumusan masalah dan kesimpulan justru menilai arsitektur mana yang unggul, kesesuaian fokus penelitian perlu dibahas bersama pembimbing.')

start=next(p for p in doc.paragraphs if p.text=='12 Jawaban sepuluh pertanyaan prioritas')
end=next(p for p in doc.paragraphs if p.text=='13 Pola menjawab saat sidang')
element=start._p.getnext()
while element is not None and element is not end._p:
    nxt=element.getnext(); element.getparent().remove(element); element=nxt

qas=[
('1. Apa novelty atau kontribusi penelitian Anda?',
'“Kontribusi penelitian saya adalah rancangan dan penerapan microservices pada SIPEKA, dengan pemisahan layanan pelaporan dan penanganan kasus, kepemilikan data per layanan, serta sinkronisasi status melalui RabbitMQ. Saya juga mengevaluasi performa implementasinya menggunakan monolith sebagai baseline. Kontribusi praktisnya adalah integrasi alur pelaporan masyarakat sampai terminasi kasus.”',
'Kontribusi belum otomatis berarti kebaruan ilmiah. Untuk menjawab novelty, tunjukkan perbedaan spesifik dengan penelitian terdahulu yang benar-benar ditelaah. Jangan mengklaim pertama, algoritma baru, atau performa lebih unggul tanpa bukti.'),
('2. Mengapa membandingkan microservices dengan monolith?',
'“Monolith saya gunakan sebagai baseline agar pengukuran performa microservices memiliki acuan pada fungsi aplikasi yang setara. Tujuan utamanya tetap mengevaluasi penerapan microservices di SIPEKA. Perbandingan ini membantu melihat konsekuensi implementasi, termasuk jika hasilnya hanya berbeda sedikit atau microservices lebih lambat.”',
'Jika penguji mempertanyakan judul, jelaskan hubungan tujuan penerapan dengan evaluasinya. Jangan menyangkal adanya perbandingan dan jangan mengubah tujuan penelitian hanya pada jawaban lisan.'),
('3. Bagaimana menjamin perbandingannya adil?',
'“Saya harus menjelaskan kesetaraan fungsi endpoint, payload, kondisi awal data, beban, durasi, timeout, dan lingkungan pengujian. Saya juga melaporkan batas resource per container dan total setiap arsitektur. Keadilan pengujian dinilai dari bukti konfigurasi dan pelaksanaannya, sehingga saya tidak menyatakan semuanya setara jika masih ada perbedaan.”',
'Pada compose dasar, dua backend microservices masing-masing dibatasi 1 CPU dan 512 MB, sedangkan monolith satu backend dengan batas yang sama. Total batas backend berbeda; database, broker, dan gateway juga perlu diperhitungkan sesuai cakupan evaluasi. File pengujian menunjuk salinan database, tetapi tidak meresetnya otomatis setelah setiap run. Gunakan konfigurasi yang benar-benar dipakai dan akui potensi faktor perancu.'),
('4. Apa kesimpulan yang benar-benar didukung oleh hasil pengujian?',
'“Jika tabel hasil menunjukkan selisih relatif kecil, kesimpulan saya adalah bahwa pada beban dan konfigurasi yang diuji, perbedaan performa yang teramati juga relatif kecil. Saya tidak menyimpulkan microservices selalu lebih cepat. Keberhasilan fungsi sistem hanya saya nyatakan untuk skenario fungsional yang telah diuji dan lulus.”',
'Isi dengan nilai aktual response time, throughput, error rate, dan variasi antarulang. Tanpa analisis yang memadai, jangan menyebut kedua arsitektur setara secara statistik atau tidak berbeda signifikan. Selisih kecil tetap sah; angka, grafik, dan log harus tetap sesuai hasil pengukuran.'),
('5. Mengapa memakai microservices jika performanya belum tentu lebih baik?',
'“Pertimbangan saya mencakup pemisahan tanggung jawab antara pelaporan dan penanganan kasus serta pengelolaan layanan secara terpisah. Keputusan arsitektur mempunyai konsekuensi biaya komunikasi dan operasional. Jika selisih performanya kecil, saya tetap melaporkan hasil tersebut dan menilai pilihan arsitektur berdasarkan kebutuhan sistem, bukan hanya kecepatan respons.”',
'Scaling independen dan isolasi kegagalan adalah potensi desain yang perlu dibuktikan melalui eksperimen tersendiri. Monolith juga dapat menjadi pilihan wajar untuk sistem kecil; jangan mengklaim microservices selalu lebih tepat.'),
('6. Mengapa hanya dua service?',
'“Saya membagi layanan berdasarkan tanggung jawab bisnis. Reporting Service menerima dan menyimpan laporan masyarakat. Case Service menangani registrasi, assessment, intervensi, monitoring, dan terminasi. Autentikasi masih digabung di Case Service untuk membatasi kompleksitas implementasi. Jumlah service mengikuti kebutuhan domain, bukan semakin banyak semakin baik.”',
'Tunjukkan route serta database masing-masing. Jika kebutuhan berkembang, pemisahan autentikasi dapat dipertimbangkan dengan memperhatikan manfaat dan tambahan biaya operasionalnya.'),
('7. Mengapa RabbitMQ diperlukan?',
'“RabbitMQ dipilih untuk menyampaikan event perubahan status dari Case Service ke Reporting Service secara asynchronous. Case Service menyimpan penanganan pada databasenya, lalu Reporting Service memperbarui status laporan setelah menerima event. Dengan begitu, Case Service tidak menulis langsung ke database Reporting Service.”',
'RabbitMQ merupakan pilihan desain, bukan syarat wajib semua microservices. HTTP langsung adalah alternatif dengan konsekuensi ketergantungan saat request berlangsung. Queue yang digunakan adalah kasus_status_updates. Pengiriman laporan awal memakai HTTP, bukan RabbitMQ.'),
('8. Apa yang terjadi jika RabbitMQ atau salah satu service mati?',
'“Dampaknya bergantung pada komponen yang mati dan apakah event sudah diterima broker. Jika RabbitMQ mati saat publish, data kasus dapat tersimpan tetapi status laporan belum tersinkron. Jika Reporting Service mati, endpoint laporan tidak tersedia; event yang sudah tersimpan di broker dapat menunggu consumer kembali. Jika Case Service mati, login dan penanganan tidak tersedia, sedangkan pelaporan dapat tetap berjalan jika Reporting Service, gateway, dan MongoDB masih tersedia.”',
'Durable queue dan persistent message bukan jaminan bahwa event yang gagal dipublikasikan telah tersimpan. Kode saat ini mencatat kegagalan publish dan belum memiliki outbox serta publisher confirms. Startup Case Service juga menjalankan seed akun/master dan worker reminder, bukan startConsumer RabbitMQ; consumer status berada pada Reporting Service. Pernyataan dampak di atas merupakan penelusuran kode, bukan hasil uji kegagalan yang sudah dilakukan.'),
('9. Apa kelemahan sistem Anda?',
'“Validasi aturan bisnis belum seluruhnya konsisten antara frontend dan backend. Pembatasan role dan urutan perubahan status juga perlu diperkuat. Sinkronisasi status masih berisiko tertinggal ketika pengiriman event gagal. Selain itu, gateway, broker, dan database masih menjadi titik ketergantungan bersama, file bukti memakai penyimpanan lokal, dan hasil evaluasi terbatas pada konfigurasi yang diuji.”',
'Pilih beberapa keterbatasan yang dapat dijelaskan beserta dampaknya. Contohnya, kronologi pendek dapat melewati validasi frontend melalui Postman; endpoint assessment/intervensi belum mengunci seluruh perubahan pada kasus arsip. Jangan menyatakan sistem sudah sepenuhnya aman atau tahan gangguan.'),
('10. Jika penelitian dilanjutkan, apa yang pertama kali akan diperbaiki?',
'“Prioritas awal saya adalah memperkuat validasi dan kontrol akses pada backend, termasuk aturan transisi status dan penguncian kasus arsip, karena berkaitan dengan keutuhan data penanganan. Setelah itu saya memperkuat pengiriman event dengan outbox, konfirmasi broker, dan mekanisme retry yang terukur. Evaluasi lanjutan dilakukan dengan pengulangan serta konfigurasi resource yang lebih terkontrol.”',
'Ini adalah usulan pengembangan, bukan fitur yang sudah ada. Untuk penelitian performa, tetapkan skenario berdasarkan pertanyaan penelitian dan laporkan seluruh hasil; jangan memilih hanya pengujian yang membuat microservices terlihat unggul.'),
]
before(end,'Gunakan contoh jawaban berikut dengan bahasa sendiri. Sesuaikan pernyataan pelaksanaan dan hasil dengan bukti dalam skripsi; jangan membaca rencana pengujian seolah-olah sudah dilakukan.')
for question,answer,note in qas:
    before(end,question,True)
    before(end,answer)
    p=before(end,'')
    p.add_run('Catatan pendalaman. ').bold=True
    p.add_run(note)

before(end,'Menanggapi selisih performa yang kecil',True)
before(end,'“Selisih kecil tetap saya laporkan apa adanya. Saya memeriksa apakah hasilnya konsisten antarpercobaan dan apakah konfigurasi memberi ruang untuk menarik kesimpulan. Jika bukti belum cukup, saya membatasi kesimpulan pada hasil yang teramati dan mengusulkan pengujian lanjutan.”')
before(end,'Koreksi kesalahan hitung boleh dilakukan dengan menyimpan data mentah dan mencatat koreksinya. Mengubah angka untuk memperbesar selisih, menghapus hasil yang tidak disukai tanpa kriteria yang sah, atau menggambar grafik yang tidak sesuai data akan merusak validitas penelitian.')
doc.core_properties.subject='Penerapan microservices pada SIPEKA dan evaluasi menggunakan baseline monolith'
doc.save(output)
print(output)
print('Questions added:',len(qas))
assert all(any(q==p.text for p in doc.paragraphs) for q,_,_ in qas)
assert 'Kontribusi penelitian: hasil perbandingan performa kedua arsitektur dalam konteks SIPEKA.' not in '\n'.join(p.text for p in doc.paragraphs)
