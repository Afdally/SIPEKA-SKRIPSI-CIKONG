from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(r'C:\Users\cikoh\.codex\attachments\ef8df16d-aac4-4ca7-8969-28a0372ec7b4\Pasted text.txt')
OUT = ROOT / 'output' / 'Panduan_Pertanyaan_Sidang_SIPEKA.docx'
text = SOURCE.read_text(encoding='utf-8-sig')
text = text[text.index('1. Pertanyaan pembuka penelitian'):]
text = text[:text.index('Pola tersebut membuat jawaban')]
text = text.replace('2. Pertanyaan paling berbahaya: metodologi pengujian', '2. Metodologi pengujian')
text = text.replace('12. Pertanyaan kritis yang paling mungkin menjebak', '12. Pertanyaan prioritas latihan')
text = text.replace('Ini sangat mungkin ditanyakan.', 'Siapkan bukti konfigurasi dan catatan pelaksanaan untuk menjawab pertanyaan ini.')
text = text.replace('Jawaban ideal:', 'Kerangka jawaban yang perlu disesuaikan dengan pelaksanaan pengujian:')
text = text.replace('“Keduanya diuji menggunakan endpoint, payload, skenario beban, timeout, dan lingkungan perangkat keras yang sama. Resource container juga dibatasi agar hasil tidak dipengaruhi perbedaan alokasi sumber daya. Variabel yang dibedakan hanya bentuk arsitekturnya.”', '“Saya membandingkan fungsi dan beban yang setara, lalu menjelaskan konfigurasi perangkat keras, database, timeout, serta resource setiap arsitektur. Saya melaporkan perbedaan konfigurasi yang masih dapat memengaruhi hasil.”\nCatatan: batas per container tidak sama dengan batas total arsitektur. Pada compose dasar, dua container backend microservices masing-masing dibatasi 1 CPU dan 512 MB, sedangkan monolith satu container dengan batas yang sama. Ini adalah batas maksimum, bukan jaminan pemakaian aktual. Jelaskan konfigurasi yang benar-benar digunakan saat eksperimen; jangan mengklaim resource total sama tanpa bukti.')
text = text.replace('- Throughput: jumlah request yang berhasil diproses per satuan waktu.', '- Throughput: jumlah request yang diproses per satuan waktu menurut definisi alat ukur. Pada laporan JMeter, throughput dapat mencakup sampel gagal; baca bersama error rate dan jangan otomatis menyebutnya throughput sukses.')
text = text.replace('- Request melewati API Gateway.', '- Periksa jalur request yang benar-benar diuji. Pada konfigurasi ini, microservices dan monolith sama-sama melewati gateway; keberadaan gateway saja tidak menjelaskan selisih performa.')
text = text.replace('- Ada komunikasi jaringan antarservice.', '- Ada biaya komunikasi jaringan antarservice pada operasi yang memang melibatkan lebih dari satu service.')
text = text.replace('- Ada overhead container dan koneksi.', '- Biaya koneksi dan pembagian resource dapat berpengaruh. Penyebab hasil harus dibuktikan dengan pengukuran, bukan hanya diduga dari nama arsitektur.')
text = re.sub(r'Ini penting karena komentar tersebut juga tercantum di .*?\n', '', text)
text = text.replace('Karena konfigurasi saat ini hanya memiliki satu instance gateway, gateway masih menjadi single point of failure.', 'Karena konfigurasi saat ini hanya memiliki satu instance gateway, gateway masih menjadi single point of failure pada jalur akses melalui gateway.')
text = text.replace('- LP: korban perempuan dewasa.', '- LP: tipe laporan perempuan. Frontend memilih tipe berdasarkan usia; model backend membuat awalan kode berdasarkan tipe_laporan yang diterima.')
text = text.replace('Pesan dibuat persistent dan queue bersifat durable. Pesan dapat menunggu sampai consumer kembali aktif.', 'Jika pesan sudah diterima dan tersimpan oleh broker yang tetap tersedia, pesan dapat menunggu sampai consumer Reporting Service aktif kembali. Queue menggunakan durable dan pesan menggunakan persistent.')
text = text.replace('“Konfigurasi tersebut meningkatkan ketahanan pesan, walaupun jaminan penuh juga bergantung pada persistence RabbitMQ dan mekanisme penanganan kegagalan yang digunakan.”', '“Durable dan persistent membantu ketahanan pesan. Namun jika RabbitMQ tidak dapat dihubungi saat publish, pesan belum tentu masuk ke antrean. Pada kode saat ini, kegagalan publish dicatat sementara perubahan kasus tetap dapat tersimpan. Belum ada transactional outbox dan publisher confirms untuk menutup celah tersebut.”')
text = text.replace('“Data di kedua service tidak harus konsisten pada milidetik yang sama. Setelah event diproses, status di Reporting Service akan menyusul status di Case Service.”', '“Data di kedua service dapat berbeda sementara. Setelah event berhasil dikirim dan diproses, status di Reporting Service menyusul status di Case Service. Jika event gagal dikirim, diperlukan pemulihan atau pengiriman ulang agar konsistensi tercapai.”')
text = text.replace('- Public GIS hanya memilih field nonidentitas.', '- Public GIS mengecualikan nama, NIK, dan telepon, tetapi data lokasi serta waktu kejadian tetap perlu dievaluasi karena dapat membantu identifikasi tidak langsung.')
text = text.replace('- ID tidak ditemukan → 404.', '- ID berformat valid tetapi dokumennya tidak ditemukan → 404. ID dengan format tidak valid dapat menghasilkan 500 pada beberapa controller saat ini.')
text = text.replace('Penguji biasanya lebih menghargai jawaban jujur disertai mitigasi daripada klaim “sistem sudah aman”.', 'Jelaskan kontrol yang sudah ada beserta keterbatasannya; hindari klaim bahwa sistem sepenuhnya aman.')
text = text.replace('Ini salah satu kelemahan nyata yang sebaiknya Anda akui.', 'Jelaskan pemeriksaan duplikat yang sudah ada serta keterbatasannya pada request bersamaan.')
text = text.replace('Pola menjawab saat sidang', '13. Pola menjawab saat sidang')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title',25), ('Subtitle',12), ('Heading 1',16), ('Heading 2',12)]:
    st=doc.styles[name]
    st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(14 if name=='Heading 1' else 8)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0]
footer.alignment=2
run=footer.add_run('SIPEKA  |  '); run.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='Panduan Pertanyaan Sidang Skripsi SIPEKA'
doc.core_properties.subject='Latihan tanya jawab dan demonstrasi sistem'
doc.add_paragraph('Panduan Pertanyaan Sidang Skripsi SIPEKA', 'Title')
doc.add_paragraph('Latihan jawaban tentang penelitian dan alur sistem', 'Subtitle')
doc.add_paragraph('Panduan ini membantu Anda menjelaskan alasan penelitian, arsitektur SIPEKA, perjalanan request, pengujian API, dan keterbatasan implementasi. Gunakan contoh jawaban untuk latihan berbicara dengan bahasa sendiri. Pertanyaan berikut merupakan bahan latihan, bukan kepastian pertanyaan yang akan diajukan penguji.')
p=doc.add_paragraph(); p.add_run('Cara belajar. ').bold=True
p.add_run('Baca satu pertanyaan, tutup jawabannya, lalu jelaskan selama 30 sampai 60 detik. Buka kode atau hasil uji untuk mendukung jawaban. Tandai pertanyaan yang belum bisa dijelaskan tanpa membaca.')
p=doc.add_paragraph(); p.add_run('Data penelitian. ').bold=True
p.add_run('Isi lembar persiapan pada bagian akhir dengan angka dan prosedur yang benar-benar digunakan. Bedakan kemampuan arsitektur, fitur yang sudah diimplementasikan, dan hasil yang sudah diuji.')

lines=text.splitlines()
i=0
while i<len(lines):
    line=lines[i].strip(); i+=1
    if not line: continue
    if re.match(r'^\d+\. (Pertanyaan|Metodologi|Pola)',line):
        heading=re.sub(r'^(\d+)\. ',r'\1 ',line).replace(':','')
        doc.add_paragraph(heading,'Heading 1'); continue
    if line.startswith('“') and line.endswith('?”'):
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
        p.add_run(line.strip('“”')).bold=True; continue
    if line.startswith('- '):
        doc.add_paragraph(line[2:],'List Bullet'); continue
    if re.match(r'^\d+\. ',line):
        doc.add_paragraph(line); continue
    if line.startswith('→') or (i<len(lines) and lines[i].strip().startswith('→')):
        chain=[line.lstrip('→ ')]
        while i<len(lines) and lines[i].strip().startswith('→'):
            chain.append(lines[i].strip().lstrip('→ ')); i+=1
        doc.add_paragraph(' → '.join(chain)); continue
    p=doc.add_paragraph(line)
    if line.endswith(':'): p.paragraph_format.keep_with_next=True

doc.add_paragraph('14 Jawaban tambahan untuk pendalaman', 'Heading 1')
qa=[
('Apa perbedaan service frontend dan microservice backend?', 'Service frontend adalah kumpulan fungsi JavaScript untuk memanggil API, misalnya laporanService.submitLaporan(). Microservice backend adalah aplikasi server tersendiri yang menjalankan logika bisnis dan mengakses database. Folder utils berisi fungsi bantuan seperti filter laporan dan notifikasi.'),
('Apa fungsi token setelah login?', 'Token JWT menjadi bukti autentikasi untuk request berikutnya. Frontend atau Postman mengirimnya melalui Authorization: Bearer <token>. Backend memeriksa tanda tangan dan masa berlakunya, lalu menggunakan identitas serta role dari token. JWT ditandatangani, tetapi isi payload tidak otomatis dienkripsi. Masa berlaku default kode adalah 3600 detik dan dapat dikonfigurasi.'),
('Apa perbedaan autentikasi dan otorisasi?', 'Autentikasi memeriksa siapa pengguna, sedangkan otorisasi menentukan tindakan yang diizinkan. Middleware auth memverifikasi JWT. Middleware role membatasi endpoint tertentu, misalnya pengelolaan pengguna untuk super_admin. Tidak semua route terproteksi sudah memiliki pembatasan role khusus.'),
('Siapa yang boleh melakukan terminasi pada implementasi saat ini?', 'Pada alur tampilan, terminasi dilakukan melalui dashboard petugas. Di backend, route penanganan memakai middleware autentikasi, tetapi endpoint terminasi belum memiliki pembatasan role khusus. Karena itu jangan menyatakan hanya petugas yang dapat memanggilnya berdasarkan kontrol backend saat ini.'),
('Apakah kasus yang selesai sudah sepenuhnya terkunci?', 'Terminasi mengubah status menjadi selesai dan arsip menjadi true. Penambahan log memeriksa arsip dan akan ditolak. Namun controller assessment dan intervensi belum memeriksa arsip secara menyeluruh, sehingga penguncian seluruh perubahan masih perlu diperkuat.'),
('Bagaimana audit trail dan koreksi assessment ditangani?', 'Activity log menyimpan catatan, tanggal, nama petugas, dan timestamp. Ini belum merupakan audit trail lengkap untuk semua perubahan. Penyimpanan assessment saat ini dapat menimpa nilai sebelumnya tanpa riwayat versi; pengembangan berikutnya dapat menambahkan histori perubahan dan aturan koreksi.'),
('Mengapa hasil Postman berbeda dari validasi formulir?', 'Postman langsung memanggil API tanpa menjalankan React. Validasi yang hanya berada di frontend dapat dilewati. Respons 422 untuk field wajib menunjukkan validasi model bekerja, tetapi belum membuktikan seluruh aturan bisnis sudah divalidasi di server.'),
('Apa yang pertama diperbaiki jika penelitian dilanjutkan?', 'Prioritas teknis yang dapat diusulkan adalah memperkuat validasi backend dan akses berbasis role, kemudian memastikan event status tidak hilang dengan outbox atau mekanisme pemulihan yang teruji. Sesuaikan prioritas akhir dengan dampak masalah dan batas penelitian.'),
]
for q,a in qa:
    p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; p.add_run(q).bold=True
    doc.add_paragraph(a)

doc.add_paragraph('15 Lembar persiapan hasil pengujian', 'Heading 1')
doc.add_paragraph('Lengkapi sebelum latihan sidang. Untuk hal yang tidak diukur atau tidak dilakukan, tulis secara jujur dan jelaskan dampaknya terhadap interpretasi hasil.')
for item in ['Judul dan rumusan masalah penelitian','Metode penelitian dan alasan pemilihannya','Endpoint dan payload pengujian','Jumlah pengguna virtual serta ramp up','Durasi atau loop dan jumlah pengulangan','Spesifikasi perangkat serta lokasi load generator','Batas resource per container dan total arsitektur','Kondisi awal database dan prosedur reset setiap run','Warm up dan urutan pelaksanaan pengujian','Response time rata rata dan P95 beserta satuan','Throughput dan error rate setiap skenario','CPU dan memori jika diukur','Kesimpulan utama beserta tabel atau grafik pendukung','Batas generalisasi dan penelitian terdahulu pembanding']:
    p=doc.add_paragraph(); p.add_run(item + ': ').bold=True; p.add_run('................................................................')
doc.add_paragraph('Contoh pola jawaban hasil: Pada skenario [beban], arsitektur [nama] menghasilkan [metrik dan nilai], dibandingkan [nilai pembanding]. Temuan ini berlaku pada [konfigurasi pengujian]. Penyebab selisih [telah didukung pengukuran atau masih berupa dugaan].')

doc.add_paragraph('16 Peta file untuk latihan demonstrasi', 'Heading 1')
for path,desc in [
('frontend/ircm-frontend/src/pages/LandingPage.jsx','Titik awal submit dan validasi frontend.'),
('frontend/ircm-frontend/src/services/laporanService.js','Pemanggilan API laporan.'),
('frontend/ircm-frontend/src/services/apiClient.js','Konfigurasi Axios dan header token.'),
('nginx/nginx.conf','Routing menuju backend dan pemblokiran endpoint internal.'),
('backend/reporting-service/app.js','Pendaftaran route dan koneksi report_db.'),
('backend/reporting-service/src/routes/laporan.js','Route publik, Multer, dan route terproteksi.'),
('backend/reporting-service/src/controllers/laporanController.js','Penyimpanan laporan dan cek status.'),
('backend/reporting-service/src/models/Laporan.js','Skema, field wajib, kode laporan, dan status awal.'),
('backend/case-service/src/controllers/kasusController.js','Registrasi sampai terminasi.'),
('backend/case-service/src/queue/publisher.js','Pengiriman event status kasus.'),
('backend/reporting-service/src/queue/consumer.js','Penerimaan event dan pembaruan status laporan.'),
('backend/case-service/src/controllers/authController.js','Login dan pembuatan JWT.'),
('docker-compose.yml dan docker-compose.pengujian.yml','Service, database, port, dan konfigurasi eksperimen.')]:
    p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; r=p.add_run(path); r.bold=True; r.font.size=Pt(10)
    doc.add_paragraph(desc)
for element in [doc.styles.element, doc._element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
print('Paragraphs:',len(doc.paragraphs))
